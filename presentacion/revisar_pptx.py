# Revisa una presentacion: formas fuera del lienzo y texto que probablemente
# desborda su caja (estimacion por ancho de caracter y altura de linea).
import sys
from pptx import Presentation

p = Presentation(sys.argv[1])
W, H = p.slide_width / 360000, p.slide_height / 360000
prob = 0
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        L = (sh.left or 0) / 360000; T = (sh.top or 0) / 360000
        w = (sh.width or 0) / 360000; h = (sh.height or 0) / 360000
        if L < -0.01 or T < -0.01 or L + w > W + 0.01 or T + h > H + 0.01:
            print('%02d  FUERA   (%.1f,%.1f) %.1fx%.1f  %s' % (i, L, T, w, h,
                  (sh.text_frame.text[:40] if sh.has_text_frame else sh.shape_type)))
            prob += 1
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        alto = 0.0
        for par in sh.text_frame.paragraphs:
            txt = ''.join(r.text for r in par.runs)
            if not txt.strip():
                continue
            pt = next((r.font.size.pt for r in par.runs if r.font.size), 11)
            # ancho medio de caracter ~ 0.50 * pt en Calibri, en cm: pt*0.0353
            cpl = max(1, int(w / (0.50 * pt * 0.0353)))
            lineas = max(1, -(-len(txt) // cpl))
            alto += lineas * pt * 0.0353 * 1.22
        if alto > h + 0.25:
            print('%02d  DESBORDA  necesita %.1f tiene %.1f cm  "%s"'
                  % (i, alto, h, sh.text_frame.text[:60].replace('\n', ' | ')))
            prob += 1
print('---', prob, 'avisos en', len(p.slides.__iter__.__self__._sldIdLst), 'diapositivas')
