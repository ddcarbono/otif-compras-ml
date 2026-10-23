# -*- coding: utf-8 -*-
"""Plantilla Uninorte medida sobre la sustentacion de Analisis Multivariado:
lienzo 25.4x14.29, barra verde, filete oro, logo, pie y tarjetas de color.
Provee los ayudantes de dibujo; el contenido lo pone cada constructor.
"""
import os, sys, copy
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SC   = os.path.dirname(os.path.abspath(__file__))
IMG  = os.path.join(SC, 'img')
LOGO = os.path.join(SC, 'media', 'logo.png')
IMG  = os.environ.get('IMG_DIR', IMG)

VERDE   = RGBColor(0x40, 0x57, 0x31)
ORO     = RGBColor(0xB0, 0x8A, 0x2D)
ORO_CL  = RGBColor(0xF2, 0xE3, 0xBE)
PALIDO  = RGBColor(0xD9, 0xE2, 0xD2)
TINTA   = RGBColor(0x1A, 0x1A, 0x1A)
GRIS    = RGBColor(0x5A, 0x5A, 0x5A)
GRIS_CL = RGBColor(0x8C, 0x8C, 0x8C)
BLANCO  = RGBColor(0xFF, 0xFF, 0xFF)
VERDECL = RGBColor(0xE8, 0xF0, 0xEB)
CREMA   = RGBColor(0xF5, 0xED, 0xE0)
NIEBLA  = RGBColor(0xEE, 0xF2, 0xEA)
LINEA   = RGBColor(0xDD, 0xD5, 0xC0)
TEJA    = RGBColor(0x8C, 0x4A, 0x2D)

FUENTE = 'Calibri'
PIE = 'Daniel Carbonó   •   Maestría en Estadística Aplicada   •   Universidad del Norte'

prs = Presentation()
prs.slide_width  = Emu(int(25.4 * 360000))
prs.slide_height = Emu(int(14.29 * 360000))
BLANK = prs.slide_layouts[6]


def nueva(fondo=BLANCO):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = fondo
    return s


def caja(s, x, y, w, h, relleno=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    sh.line.fill.background()
    sh.shadow.inherit = False
    if relleno is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = relleno
    sh.text_frame.word_wrap = True
    return sh


def texto(s, x, y, w, h, lineas, tam=11, color=TINTA, negrita=False,
          espacio=None, anclaje=MSO_ANCHOR.TOP, alineacion=PP_ALIGN.LEFT):
    """lineas: str, o lista de str, o lista de (str, dict) con overrides."""
    if isinstance(lineas, str):
        lineas = [lineas]
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anclaje
    for i, ln in enumerate(lineas):
        op = {}
        if isinstance(ln, tuple):
            ln, op = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = op.get('alineacion', alineacion)
        if espacio:
            p.space_after = Pt(espacio)
        if op.get('antes'):
            p.space_before = Pt(op['antes'])
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name = op.get('fuente', FUENTE)
        f.size = Pt(op.get('tam', tam))
        f.bold = op.get('negrita', negrita)
        f.color.rgb = op.get('color', color)
    return tb


def rico(s, x, y, w, h, partes, tam=11, color=TINTA):
    """Un solo parrafo con tramos de formato distinto: [(txt, negrita, color)]."""
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    for txt, neg, col in partes:
        r = p.add_run()
        r.text = txt
        r.font.name = FUENTE
        r.font.size = Pt(tam)
        r.font.bold = neg
        r.font.color.rgb = col or color
    return tb


def filete(s, x, y, w=1.3, color=ORO, alto=0.1):
    return caja(s, x, y, w, alto, color)


def logo(s, x, y, w=1.3, h=1.2):
    caja(s, x, y, w, h, BLANCO)
    if os.path.exists(LOGO):
        m = 0.12 * w
        s.shapes.add_picture(LOGO, Cm(x + m), Cm(y + h * 0.17),
                             width=Cm(w - 2 * m))


def cabecera(s, titulo, kicker=None):
    caja(s, 0, 0, 25.4, 1.6, VERDE)
    caja(s, 0, 1.6, 25.4, 0.1, ORO)
    logo(s, 0.4, 0.2)
    texto(s, 2.2, 0.35, 22.7, 1.0, titulo, tam=20, color=BLANCO, negrita=True,
          anclaje=MSO_ANCHOR.MIDDLE)
    if kicker:
        texto(s, 1.3, 2.1, 20.3, 0.7, kicker, tam=12, color=ORO, negrita=True)
        filete(s, 1.3, 2.9)


def vineta(s, x, y, ancho, txt, tam=13, alto=1.2, color=VERDE, negrita=True):
    caja(s, x, y + 0.16, 0.2, 0.2, VERDE)
    return texto(s, x + 0.6, y, ancho, alto, txt, tam=tam, color=color, negrita=negrita)


def tarjeta(s, x, y, w, h, titulo, cuerpo, relleno=VERDECL, tam_t=13, tam_c=11,
            filete_arriba=False):
    caja(s, x, y, w, h, relleno)
    if filete_arriba:
        filete(s, x + 0.25, y, w - 0.5)
    texto(s, x + 0.45, y + 0.3, w - 0.9, 0.8, titulo, tam=tam_t, color=VERDE, negrita=True)
    if cuerpo:
        texto(s, x + 0.45, y + 1.02, w - 0.9, h - 1.15, cuerpo, tam=tam_c, color=TINTA)


def cifra(s, x, y, w, h, numero, etiqueta, relleno=VERDECL):
    caja(s, x, y, w, h, relleno)
    filete(s, x + 0.25, y, w - 0.5)
    texto(s, x, y + 0.35, w, 1.3, numero, tam=34, color=VERDE, negrita=True,
          alineacion=PP_ALIGN.CENTER)
    texto(s, x, y + 1.75, w, 0.8, etiqueta, tam=10.5, color=GRIS,
          alineacion=PP_ALIGN.CENTER)


def banda(s, x, y, w, h, partes, relleno=VERDECL, tam=13, color=VERDE):
    caja(s, x, y, w, h, relleno)
    rico(s, x + 0.5, y + 0.18, w - 1.0, h - 0.3, partes, tam=tam, color=color)


def figura(s, ruta, x, y, w=None, h=None):
    kw = {}
    if w: kw['width'] = Cm(w)
    if h: kw['height'] = Cm(h)
    return s.shapes.add_picture(os.path.join(IMG, ruta), Cm(x), Cm(y), **kw)


def tabla(s, x, y, w, h, filas, anchos=None, tam=10.5, tam_cab=10.5):
    nf, nc = len(filas), len(filas[0])
    g = s.shapes.add_table(nf, nc, Cm(x), Cm(y), Cm(w), Cm(h)).table
    if anchos:
        tot = sum(anchos)
        for j, a in enumerate(anchos):
            g.columns[j].width = Emu(int(Cm(w) * a / tot))
    for i, fila in enumerate(filas):
        g.rows[i].height = Cm(h / nf)
        for j, val in enumerate(fila):
            cel = g.cell(i, j)
            cel.margin_left = cel.margin_right = Cm(0.16)
            cel.margin_top = cel.margin_bottom = Cm(0.04)
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            cel.fill.solid()
            cel.fill.fore_color.rgb = (VERDE if i == 0
                                       else (BLANCO if i % 2 else NIEBLA))
            tf = cel.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name = FUENTE
            r.font.size = Pt(tam_cab if i == 0 else tam)
            r.font.bold = (i == 0 or j == 0)
            r.font.color.rgb = BLANCO if i == 0 else TINTA
    return g


